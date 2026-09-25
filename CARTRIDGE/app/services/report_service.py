import io
import os
import calendar
from datetime import datetime
from typing import Optional, Dict, Any, List
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import or_, and_, desc, func

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from app.models import Cartridge, CartridgeStatus, HistoryLog, Branch, AppUser, Batch
from app.services.settings_service import SettingsService

STATUS_NAMES_RU = {
    CartridgeStatus.IN_USE: "В работе",
    CartridgeStatus.PENDING_VENDOR: "Ожидает заправщика",
    CartridgeStatus.AT_VENDOR: "На заправке",
    CartridgeStatus.READY_FOR_PICKUP: "Готов к выдаче"
}

MONTHS_RU = [
    "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
    "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"
]

_PDF_FONT_REGULAR = None
_PDF_FONT_BOLD = None


def _get_pdf_fonts():
    global _PDF_FONT_REGULAR, _PDF_FONT_BOLD
    if _PDF_FONT_REGULAR:
        return _PDF_FONT_REGULAR, _PDF_FONT_BOLD

    font_candidates = [
        # Windows
        ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
        ("C:/Windows/Fonts/calibri.ttf", "C:/Windows/Fonts/calibrib.ttf"),
        # Linux / Docker
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf", "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"),
        ("/usr/share/fonts/truetype/freefont/FreeSans.ttf", "/usr/share/fonts/truetype/freefont/FreeSansBold.ttf"),
    ]
    for regular_path, bold_path in font_candidates:
        if os.path.exists(regular_path):
            try:
                pdfmetrics.registerFont(TTFont("AppReportFont", regular_path))
                if os.path.exists(bold_path):
                    pdfmetrics.registerFont(TTFont("AppReportFont-Bold", bold_path))
                    _PDF_FONT_BOLD = "AppReportFont-Bold"
                else:
                    _PDF_FONT_BOLD = "AppReportFont"
                _PDF_FONT_REGULAR = "AppReportFont"
                return _PDF_FONT_REGULAR, _PDF_FONT_BOLD
            except Exception:
                pass

    _PDF_FONT_REGULAR = "Helvetica"
    _PDF_FONT_BOLD = "Helvetica-Bold"
    return _PDF_FONT_REGULAR, _PDF_FONT_BOLD


def is_refill_action(action: Optional[str]) -> bool:
    """Определяет, относится ли действие в журнале к завершенной заправке картриджа."""
    if not action:
        return False
    act = action.strip().lower()
    # Возврат с заправки от поставщика — каноническое завершение цикла заправки
    if "возврат" in act:
        return True
    # Устаревшие или ручные записи "Приемка" (для совместимости)
    if "приемка" in act:
        return True
    return False


def get_refills_map(
    db: Session,
    cartridge_ids: List[int],
    dt_start: Optional[datetime] = None,
    dt_end: Optional[datetime] = None
) -> Dict[int, int]:
    """Подсчитывает количество заправок для списка картриджей за указанный период или всё время."""
    if not cartridge_ids:
        return {}

    query = db.query(HistoryLog.cartridge_id, HistoryLog.action).filter(
        HistoryLog.cartridge_id.in_(cartridge_ids),
        or_(
            HistoryLog.action.like("%Возврат%"),
            HistoryLog.action.like("%возврат%"),
            HistoryLog.action.like("%Приемка%"),
            HistoryLog.action.like("%приемка%")
        )
    )
    if dt_start:
        query = query.filter(HistoryLog.timestamp >= dt_start)
    if dt_end:
        query = query.filter(HistoryLog.timestamp <= dt_end)

    refills_count: Dict[int, int] = {}
    for cid, act in query.all():
        if is_refill_action(act):
            refills_count[cid] = refills_count.get(cid, 0) + 1
    return refills_count


class ReportService:
    @classmethod
    def get_report_data(
        cls,
        db: Session,
        current_user: AppUser,
        report_type: str = "all",
        branch_id: Optional[int] = None,
        year: Optional[int] = None,
        month: Optional[int] = None,
        date_from: Optional[str] = None,
        date_to: Optional[str] = None
    ) -> Dict[str, Any]:
        """
        Формирует структурированные данные отчета.
        Проверяет права доступа:
        - суперпользователь: любой филиал или все сразу
        - администратор/оператор с branch_id: строго только свой филиал
        - администратор/оператор без branch_id: любой филиал или все сразу
        """
        effective_branch_id = branch_id
        is_branch_locked = False
        if current_user:
            if current_user.role == "user":
                raise PermissionError("Доступ к формированию отчетов запрещен для вашей роли.")

            # Ограничение филиала по роли
            if current_user.role in ("admin", "operator") and current_user.branch_id:
                effective_branch_id = current_user.branch_id
                is_branch_locked = True

        branch_obj = None
        branch_name = "Все филиалы"
        if effective_branch_id:
            branch_obj = db.query(Branch).filter(Branch.id == effective_branch_id).first()
            if branch_obj:
                branch_name = branch_obj.name

        now = datetime.utcnow()
        dt_start = None
        dt_end = None
        period_title = "Все время"

        # 1. Определение временных рамок
        if report_type == "models":
            period_title = "Статистика по моделям картриджей"
        elif report_type == "year":
            y = year or now.year
            dt_start = datetime(y, 1, 1, 0, 0, 0)
            dt_end = datetime(y, 12, 31, 23, 59, 59)
            period_title = f"{y} год"
        elif report_type == "month":
            y = year or now.year
            m = month or now.month
            _, last_day = calendar.monthrange(y, m)
            dt_start = datetime(y, m, 1, 0, 0, 0)
            dt_end = datetime(y, m, last_day, 23, 59, 59)
            period_title = f"{MONTHS_RU[m - 1]} {y} г."
        elif report_type == "custom":
            try:
                dt_start = datetime.strptime(date_from, "%Y-%m-%d") if date_from else datetime(now.year, 1, 1)
            except Exception:
                dt_start = datetime(now.year, 1, 1)
            try:
                dt_end = datetime.strptime(date_to, "%Y-%m-%d").replace(hour=23, minute=59, second=59) if date_to else now
            except Exception:
                dt_end = now
            period_title = f"с {dt_start.strftime('%d.%m.%Y')} по {dt_end.strftime('%d.%m.%Y')}"
        elif report_type == "history":
            if date_from or date_to:
                try:
                    dt_start = datetime.strptime(date_from, "%Y-%m-%d") if date_from else datetime(now.year, 1, 1)
                except Exception:
                    dt_start = datetime(now.year, 1, 1)
                try:
                    dt_end = datetime.strptime(date_to, "%Y-%m-%d").replace(hour=23, minute=59, second=59) if date_to else now
                except Exception:
                    dt_end = now
                period_title = f"с {dt_start.strftime('%d.%m.%Y')} по {dt_end.strftime('%d.%m.%Y')}"
            elif month:
                y = year or now.year
                m = month
                _, last_day = calendar.monthrange(y, m)
                dt_start = datetime(y, m, 1, 0, 0, 0)
                dt_end = datetime(y, m, last_day, 23, 59, 59)
                period_title = f"{MONTHS_RU[m - 1]} {y} г."
            elif year:
                y = year
                dt_start = datetime(y, 1, 1, 0, 0, 0)
                dt_end = datetime(y, 12, 31, 23, 59, 59)
                period_title = f"{y} год"
            else:
                period_title = "Все время"
        else:
            report_type = "all"
            period_title = "Все картриджи за весь период"

        # 2. Базовый запрос картриджей филиала
        cart_query = db.query(Cartridge).options(
            joinedload(Cartridge.branch),
            joinedload(Cartridge.current_user)
        )
        if effective_branch_id:
            cart_query = cart_query.filter(Cartridge.branch_id == effective_branch_id)
        all_branch_cartridges = cart_query.all()
        cartridge_map = {c.id: c for c in all_branch_cartridges}

        if report_type == "models":
            # ОТЧЕТ ПО МОДЕЛЯМ КАРТРИДЖЕЙ
            models_map = {}
            for c in all_branch_cartridges:
                m_name = (c.model or "Не указана").strip()
                if m_name not in models_map:
                    models_map[m_name] = {
                        "model": m_name,
                        "total": 0,
                        "in_use": 0,
                        "pending_vendor": 0,
                        "at_vendor": 0,
                        "ready_for_pickup": 0,
                        "refill_count": 0,
                        "branch_name": branch_name
                    }
                entry = models_map[m_name]
                entry["total"] += 1
                if c.status == CartridgeStatus.IN_USE:
                    entry["in_use"] += 1
                elif c.status == CartridgeStatus.PENDING_VENDOR:
                    entry["pending_vendor"] += 1
                elif c.status == CartridgeStatus.AT_VENDOR:
                    entry["at_vendor"] += 1
                elif c.status == CartridgeStatus.READY_FOR_PICKUP:
                    entry["ready_for_pickup"] += 1

            if all_branch_cartridges:
                cart_ids = [c.id for c in all_branch_cartridges]
                refills_by_cart = get_refills_map(db, cart_ids)
                for c in all_branch_cartridges:
                    m_name = (c.model or "Не указана").strip()
                    if m_name in models_map:
                        models_map[m_name]["refill_count"] += refills_by_cart.get(c.id, 0)

            total_carts = len(all_branch_cartridges)
            models_list = list(models_map.values())
            models_list.sort(key=lambda x: x["total"], reverse=True)

            for m in models_list:
                m["percentage"] = round((m["total"] / total_carts * 100), 1) if total_carts > 0 else 0
                m["percentage_label"] = f"{m['percentage']}%"

            summary = {
                "total_models": len(models_list),
                "total_cartridges": total_carts,
                "in_use": sum(m["in_use"] for m in models_list),
                "pending_vendor": sum(m["pending_vendor"] for m in models_list),
                "at_vendor": sum(m["at_vendor"] for m in models_list),
                "ready_for_pickup": sum(m["ready_for_pickup"] for m in models_list),
                "total_refills": sum(m["refill_count"] for m in models_list),
                "top_model": models_list[0]["model"] if models_list else "—"
            }

            return {
                "report_type": "models",
                "period_title": period_title,
                "branch_name": branch_name,
                "branch_id": effective_branch_id,
                "is_branch_locked": is_branch_locked,
                "generated_at": now.strftime("%d.%m.%Y %H:%M"),
                "generated_by": current_user.full_name if current_user else "Система",
                "summary": summary,
                "items": models_list
            }

        elif report_type == "all":
            # ОБЩИЙ ОТЧЕТ: Полный реестр картриджей
            cart_ids = [c.id for c in all_branch_cartridges]
            refills_by_cart = get_refills_map(db, cart_ids)

            # Получаем последнее действие для каждого картриджа
            last_logs: Dict[int, HistoryLog] = {}
            if cart_ids:
                all_logs = db.query(HistoryLog).filter(
                    HistoryLog.cartridge_id.in_(cart_ids)
                ).order_by(desc(HistoryLog.timestamp)).all()
                for l in all_logs:
                    if l.cartridge_id not in last_logs:
                        last_logs[l.cartridge_id] = l

            cartridges_data = []
            for c in all_branch_cartridges:
                user_display = "—"
                user_phone = "—"
                if c.current_user:
                    user_display = c.current_user.display_name
                    user_phone = c.current_user.phone or "—"
                elif c.current_user_id:
                    user_display = c.current_user_id

                refill_count = refills_by_cart.get(c.id, 0)
                last_log = last_logs.get(c.id)

                c_cond = getattr(c, "condition", "working") or "working"
                cartridges_data.append({
                    "id": c.id,
                    "marker_label": c.marker_label,
                    "qr_code": c.qr_code or "—",
                    "model": c.model,
                    "cabinet": c.cabinet,
                    "status": c.status.value if hasattr(c.status, "value") else str(c.status),
                    "status_label": STATUS_NAMES_RU.get(c.status, str(c.status)),
                    "condition": c_cond,
                    "condition_label": "В рабочем состоянии" if c_cond == "working" else "В нерабочем состоянии",
                    "branch_name": c.branch.name if c.branch else "Не указан",
                    "user_name": user_display,
                    "user_phone": user_phone,
                    "refill_count": refill_count,
                    "last_action": last_log.action if last_log else "Создание",
                    "last_action_date": last_log.timestamp.strftime("%d.%m.%Y %H:%M") if last_log else c.updated_at.strftime("%d.%m.%Y %H:%M") if c.updated_at else "—",
                    "notes": c.notes or ""
                })

            summary = {
                "total": len(cartridges_data),
                "in_use": sum(1 for c in cartridges_data if c["status"] == CartridgeStatus.IN_USE.value),
                "pending_vendor": sum(1 for c in cartridges_data if c["status"] == CartridgeStatus.PENDING_VENDOR.value),
                "at_vendor": sum(1 for c in cartridges_data if c["status"] == CartridgeStatus.AT_VENDOR.value),
                "ready_for_pickup": sum(1 for c in cartridges_data if c["status"] == CartridgeStatus.READY_FOR_PICKUP.value),
                "total_refills": sum(c["refill_count"] for c in cartridges_data),
                "working_count": sum(1 for c in cartridges_data if c["condition"] == "working"),
                "broken_count": sum(1 for c in cartridges_data if c["condition"] == "broken")
            }

            return {
                "report_type": report_type,
                "period_title": period_title,
                "branch_name": branch_name,
                "branch_id": effective_branch_id,
                "is_branch_locked": is_branch_locked,
                "generated_at": now.strftime("%d.%m.%Y %H:%M"),
                "generated_by": current_user.full_name if current_user else "Система",
                "summary": summary,
                "items": cartridges_data
            }

        elif report_type in ("year", "month", "custom"):
            # ОТЧЕТ ЗА ИНТЕРВАЛ (год / месяц / интервал):
            # Отображает текущее состояние картриджей, участвовавших в обороте за данный период
            period_log_query = db.query(HistoryLog).join(Cartridge).filter(
                HistoryLog.timestamp >= dt_start,
                HistoryLog.timestamp <= dt_end
            )
            if effective_branch_id:
                period_log_query = period_log_query.filter(Cartridge.branch_id == effective_branch_id)

            period_logs = period_log_query.order_by(desc(HistoryLog.timestamp)).all()
            period_cart_ids = set(l.cartridge_id for l in period_logs)

            # Также учитываем картриджи, обновленные или созданные в этот период
            for c in all_branch_cartridges:
                if c.updated_at and dt_start <= c.updated_at <= dt_end:
                    period_cart_ids.add(c.id)

            active_cartridges = [cartridge_map[cid] for cid in period_cart_ids if cid in cartridge_map]
            active_cartridges.sort(key=lambda c: c.marker_label)

            # Заправки, выполненные именно в данный интервал
            refills_in_period = get_refills_map(db, [c.id for c in active_cartridges], dt_start=dt_start, dt_end=dt_end)

            # Последнее действие по картриджу в рамках выбранного периода
            last_logs_period: Dict[int, HistoryLog] = {}
            for l in period_logs:
                if l.cartridge_id not in last_logs_period:
                    last_logs_period[l.cartridge_id] = l

            cartridges_data = []
            for c in active_cartridges:
                user_display = "—"
                user_phone = "—"
                if c.current_user:
                    user_display = c.current_user.display_name
                    user_phone = c.current_user.phone or "—"
                elif c.current_user_id:
                    user_display = c.current_user_id

                last_l = last_logs_period.get(c.id)

                cartridges_data.append({
                    "id": c.id,
                    "marker_label": c.marker_label,
                    "qr_code": c.qr_code or "—",
                    "model": c.model,
                    "cabinet": c.cabinet,
                    "status": c.status.value if hasattr(c.status, "value") else str(c.status),
                    "status_label": STATUS_NAMES_RU.get(c.status, str(c.status)),
                    "branch_name": c.branch.name if c.branch else "Не указан",
                    "user_name": user_display,
                    "user_phone": user_phone,
                    "refill_count": refills_in_period.get(c.id, 0),
                    "last_action": last_l.action if last_l else "В обороте",
                    "last_action_date": last_l.timestamp.strftime("%d.%m.%Y %H:%M") if last_l else c.updated_at.strftime("%d.%m.%Y %H:%M") if c.updated_at else "—",
                    "notes": c.notes or ""
                })

            summary = {
                "total": len(cartridges_data),
                "in_use": sum(1 for c in cartridges_data if c["status"] == CartridgeStatus.IN_USE.value),
                "pending_vendor": sum(1 for c in cartridges_data if c["status"] == CartridgeStatus.PENDING_VENDOR.value),
                "at_vendor": sum(1 for c in cartridges_data if c["status"] == CartridgeStatus.AT_VENDOR.value),
                "ready_for_pickup": sum(1 for c in cartridges_data if c["status"] == CartridgeStatus.READY_FOR_PICKUP.value),
                "total_refills": sum(c["refill_count"] for c in cartridges_data)
            }

            return {
                "report_type": report_type,
                "period_title": period_title,
                "branch_name": branch_name,
                "branch_id": effective_branch_id,
                "is_branch_locked": is_branch_locked,
                "date_from": dt_start.strftime("%Y-%m-%d") if dt_start else None,
                "date_to": dt_end.strftime("%Y-%m-%d") if dt_end else None,
                "generated_at": now.strftime("%d.%m.%Y %H:%M"),
                "generated_by": current_user.full_name if current_user else "Система",
                "summary": summary,
                "items": cartridges_data
            }

        else:
            # report_type == "history":
            # ИСТОРИЯ ДВИЖЕНИЯ КАРТРИДЖЕЙ (полный журнал всех операций)
            log_query = db.query(HistoryLog).join(Cartridge)
            if dt_start:
                log_query = log_query.filter(HistoryLog.timestamp >= dt_start)
            if dt_end:
                log_query = log_query.filter(HistoryLog.timestamp <= dt_end)
            if effective_branch_id:
                log_query = log_query.filter(Cartridge.branch_id == effective_branch_id)

            logs = log_query.order_by(desc(HistoryLog.timestamp)).all()

            accepted_count = sum(1 for l in logs if "прием" in l.action.lower() or "принят" in l.action.lower())
            sent_vendor_count = sum(1 for l in logs if "передач" in l.action.lower() or "поставщик" in l.action.lower())
            returned_vendor_count = sum(1 for l in logs if "возврат" in l.action.lower())
            issued_count = sum(1 for l in logs if "выдач" in l.action.lower())
            wa_notifications = sum(1 for l in logs if "whatsapp" in l.action.lower())

            batch_q = db.query(Batch)
            if dt_start:
                batch_q = batch_q.filter(Batch.created_at >= dt_start)
            if dt_end:
                batch_q = batch_q.filter(Batch.created_at <= dt_end)
            if effective_branch_id:
                batch_q = batch_q.filter(Batch.branch_id == effective_branch_id)
            batches_count = batch_q.count()

            unique_cart_ids = set(l.cartridge_id for l in logs)

            log_items = []
            for l in logs:
                c = cartridge_map.get(l.cartridge_id) or l.cartridge
                b_name = "—"
                if c and c.branch:
                    b_name = c.branch.name
                elif branch_obj:
                    b_name = branch_obj.name

                log_items.append({
                    "id": l.id,
                    "timestamp": l.timestamp.strftime("%d.%m.%Y %H:%M"),
                    "cartridge_marker": c.marker_label if c else f"ID {l.cartridge_id}",
                    "cartridge_model": c.model if c else "—",
                    "branch_name": b_name,
                    "action": l.action,
                    "user_name": l.user_name or "—",
                    "details": l.details or ""
                })

            summary = {
                "total_operations": len(logs),
                "unique_cartridges": len(unique_cart_ids),
                "accepted_count": accepted_count,
                "sent_vendor_count": sent_vendor_count,
                "returned_vendor_count": returned_vendor_count,
                "issued_count": issued_count,
                "wa_notifications": wa_notifications,
                "batches_count": batches_count
            }

            return {
                "report_type": "history",
                "period_title": period_title,
                "branch_name": branch_name,
                "branch_id": effective_branch_id,
                "is_branch_locked": is_branch_locked,
                "date_from": dt_start.strftime("%Y-%m-%d") if dt_start else None,
                "date_to": dt_end.strftime("%Y-%m-%d") if dt_end else None,
                "generated_at": now.strftime("%d.%m.%Y %H:%M"),
                "generated_by": current_user.full_name if current_user else "Система",
                "summary": summary,
                "items": log_items
            }

    @classmethod
    def generate_excel(cls, report: Dict[str, Any], org_name: str = "") -> io.BytesIO:
        """
        Создает стилизованный Excel-файл (.xlsx) с автоформатированием колонок и сводкой.
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "Отчет"
        ws.views.sheetView[0].showGridLines = True

        # Стили
        font_title = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
        font_sub = Font(name="Calibri", size=10, italic=True, color="475569")
        font_meta = Font(name="Calibri", size=10, bold=True, color="1E293B")
        font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        font_data = Font(name="Calibri", size=10, color="0F172A")
        font_kpi_num = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
        font_kpi_lbl = Font(name="Calibri", size=9, color="64748B")

        fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        fill_kpi = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

        border_thin = Side(border_style="thin", color="CBD5E1")
        cell_border = Border(top=border_thin, left=border_thin, right=border_thin, bottom=border_thin)

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")

        company = org_name or "Cartridge Tracker"
        rep_type = report.get("report_type", "all")

        if rep_type == "models":
            report_title = "ОТЧЕТ ПО МОДЕЛЯМ КАРТРИДЖЕЙ"
        elif rep_type == "history":
            report_title = f"ИСТОРИЯ ДВИЖЕНИЯ КАРТРИДЖЕЙ ({report['period_title'].upper()})"
        elif rep_type == "all":
            report_title = "ОБЩИЙ РЕЕСТР КАРТРИДЖЕЙ"
        else:
            report_title = f"ОТЧЕТ ПО СОСТОЯНИЮ КАРТРИДЖЕЙ ({report['period_title'].upper()})"

        ws.cell(row=1, column=1, value=company.upper()).font = font_sub
        ws.cell(row=2, column=1, value=report_title).font = font_title
        ws.cell(
            row=3,
            column=1,
            value=f"Филиал: {report['branch_name']}   |   Период: {report['period_title']}   |   Сформирован: {report['generated_at']} ({report['generated_by']})"
        ).font = font_meta

        current_row = 5
        summary = report.get("summary", {})

        # KPI блок
        if rep_type == "models":
            kpis = [
                ("Всего моделей", summary.get("total_models", 0)),
                ("Всего картриджей", summary.get("total_cartridges", 0)),
                ("В работе (у коллег)", summary.get("in_use", 0)),
                ("Ожидает заправщика", summary.get("pending_vendor", 0)),
                ("На заправке", summary.get("at_vendor", 0)),
                ("Готовы к выдаче", summary.get("ready_for_pickup", 0)),
                ("Всего заправок", summary.get("total_refills", 0)),
            ]
        elif rep_type == "history":
            kpis = [
                ("Операций за период", summary.get("total_operations", 0)),
                ("Картриджей в обороте", summary.get("unique_cartridges", 0)),
                ("Принято в IT", summary.get("accepted_count", 0)),
                ("Отправлено поставщику", summary.get("sent_vendor_count", 0)),
                ("Возвращено / Заправлено", summary.get("returned_vendor_count", 0)),
                ("Выдано в работу", summary.get("issued_count", 0)),
                ("Актов передачи", summary.get("batches_count", 0)),
            ]
        else:
            kpis = [
                ("Всего картриджей", summary.get("total", 0)),
                ("В работе (у коллег)", summary.get("in_use", 0)),
                ("Ожидает заправщика", summary.get("pending_vendor", 0)),
                ("На заправке", summary.get("at_vendor", 0)),
                ("Готовы к выдаче", summary.get("ready_for_pickup", 0)),
                ("Заправок за период", summary.get("total_refills", 0)),
            ]

        col_idx = 1
        for lbl, val in kpis:
            c_val = ws.cell(row=current_row, column=col_idx, value=val)
            c_val.font = font_kpi_num
            c_val.alignment = align_center
            c_val.fill = fill_kpi
            c_val.border = cell_border

            c_lbl = ws.cell(row=current_row + 1, column=col_idx, value=lbl)
            c_lbl.font = font_kpi_lbl
            c_lbl.alignment = align_center
            c_lbl.fill = fill_kpi
            c_lbl.border = cell_border
            col_idx += 1

        current_row += 3

        # Таблица данных
        if rep_type == "models":
            headers = [
                "№", "Модель картриджа", "Всего шт.", "Доля парка",
                "В работе", "Ожидает заправщика", "На заправке",
                "Готов к выдаче", "Всего заправок"
            ]
            fields = [
                "model", "total", "percentage_label",
                "in_use", "pending_vendor", "at_vendor",
                "ready_for_pickup", "refill_count"
            ]
        elif rep_type == "history":
            headers = [
                "№", "Дата / Время", "Картридж", "Модель", "Филиал",
                "Операция", "Исполнитель / Сотрудник", "Детали операции"
            ]
            fields = [
                "timestamp", "cartridge_marker", "cartridge_model", "branch_name",
                "action", "user_name", "details"
            ]
        else:
            headers = [
                "№", "Метка", "QR-код", "Модель", "Кабинет", "Статус",
                "Филиал", "Текущий владелец", "Телефон", "Кол-во заправок",
                "Последнее действие", "Дата действия", "Примечания"
            ]
            fields = [
                "marker_label", "qr_code", "model", "cabinet", "status_label",
                "branch_name", "user_name", "user_phone", "refill_count",
                "last_action", "last_action_date", "notes"
            ]

        for c_idx, h_text in enumerate(headers, 1):
            cell = ws.cell(row=current_row, column=c_idx, value=h_text)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = cell_border

        ws.row_dimensions[current_row].height = 25
        current_row += 1

        items = report.get("items", [])
        for idx, item in enumerate(items, 1):
            is_even = (idx % 2 == 0)
            row_fill = fill_zebra if is_even else None

            cell_num = ws.cell(row=current_row, column=1, value=idx)
            cell_num.font = font_data
            cell_num.alignment = align_center
            cell_num.border = cell_border
            if row_fill:
                cell_num.fill = row_fill

            for c_idx, fld in enumerate(fields, 2):
                val = item.get(fld, "")
                cell = ws.cell(row=current_row, column=c_idx, value=val)
                cell.font = font_data
                cell.border = cell_border
                if row_fill:
                    cell.fill = row_fill

                if fld in ("timestamp", "last_action_date", "refill_count", "cabinet", "status_label"):
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left

            ws.row_dimensions[current_row].height = 20
            current_row += 1

        for col in ws.columns:
            max_len = 0
            for cell in col:
                if cell.row < 5:
                    continue
                v_str = str(cell.value or "")
                if len(v_str) > max_len:
                    max_len = len(v_str)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(min(max_len + 3, 50), 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @classmethod
    def generate_pdf(cls, report: Dict[str, Any], org_name: str = "") -> io.BytesIO:
        """
        Создает стилизованный PDF-файл в альбомной ориентации (A4 landscape)
        с автоформатированием, поддержкой кириллицы, KPI-сводкой и таблицей данных.
        """
        font_regular, font_bold = _get_pdf_fonts()

        stream = io.BytesIO()
        doc = SimpleDocTemplate(
            stream,
            pagesize=landscape(A4),
            leftMargin=20,
            rightMargin=20,
            topMargin=20,
            bottomMargin=20
        )

        company = org_name or "Cartridge Tracker"
        rep_type = report.get("report_type", "all")

        if rep_type == "models":
            report_title = "ОТЧЕТ ПО МОДЕЛЯМ КАРТРИДЖЕЙ"
        elif rep_type == "history":
            report_title = f"ИСТОРИЯ ДВИЖЕНИЯ КАРТРИДЖЕЙ ({report['period_title'].upper()})"
        elif rep_type == "all":
            report_title = "ОБЩИЙ РЕЕСТР КАРТРИДЖЕЙ"
        else:
            report_title = f"ОТЧЕТ ПО СОСТОЯНИЮ КАРТРИДЖЕЙ ({report['period_title'].upper()})"

        styles = getSampleStyleSheet()
        p_title = ParagraphStyle('RepTitle', fontName=font_bold, fontSize=12, leading=15, textColor=colors.HexColor('#1E3A8A'))
        p_sub = ParagraphStyle('RepSub', fontName=font_bold, fontSize=7.5, leading=10, textColor=colors.HexColor('#475569'))
        p_meta = ParagraphStyle('RepMeta', fontName=font_bold, fontSize=8, leading=11, textColor=colors.HexColor('#1E293B'))

        p_kpi_val = ParagraphStyle('KpiVal', fontName=font_bold, fontSize=11, leading=13, alignment=1, textColor=colors.HexColor('#1E3A8A'))
        p_kpi_lbl = ParagraphStyle('KpiLbl', fontName=font_bold, fontSize=7, leading=9, alignment=1, textColor=colors.HexColor('#64748B'))

        p_head = ParagraphStyle('TH', fontName=font_bold, fontSize=7.5, leading=9, alignment=1, textColor=colors.white)
        p_cell = ParagraphStyle('TD', fontName=font_regular, fontSize=7, leading=9, textColor=colors.HexColor('#0F172A'))
        p_cell_c = ParagraphStyle('TDC', fontName=font_regular, fontSize=7, leading=9, alignment=1, textColor=colors.HexColor('#0F172A'))

        elements = []

        # Шапка
        elements.append(Paragraph(company.upper(), p_sub))
        elements.append(Paragraph(report_title, p_title))
        elements.append(Paragraph(
            f"Филиал: {report['branch_name']}   |   Период: {report['period_title']}   |   Сформирован: {report['generated_at']} ({report['generated_by']})",
            p_meta
        ))
        elements.append(Spacer(1, 8))

        # KPI сводка
        summary = report.get("summary", {})
        if rep_type == "models":
            kpi_items = [
                ("Всего моделей", summary.get("total_models", 0)),
                ("Всего картриджей", summary.get("total_cartridges", 0)),
                ("В работе", summary.get("in_use", 0)),
                ("Ожидает заправщика", summary.get("pending_vendor", 0)),
                ("На заправке", summary.get("at_vendor", 0)),
                ("Готовы к выдаче", summary.get("ready_for_pickup", 0)),
                ("Всего заправок", summary.get("total_refills", 0)),
            ]
        elif rep_type == "history":
            kpi_items = [
                ("Всего операций", summary.get("total_operations", 0)),
                ("Картриджей в обороте", summary.get("unique_cartridges", 0)),
                ("Принято в IT", summary.get("accepted_count", 0)),
                ("Сдано поставщику", summary.get("sent_vendor_count", 0)),
                ("Возврат с заправки", summary.get("returned_vendor_count", 0)),
                ("Выдано в работу", summary.get("issued_count", 0)),
                ("Актов передачи", summary.get("batches_count", 0)),
            ]
        else:
            kpi_items = [
                ("Всего картриджей", summary.get("total", 0)),
                ("В работе", summary.get("in_use", 0)),
                ("Ожидает заправщика", summary.get("pending_vendor", 0)),
                ("На заправке", summary.get("at_vendor", 0)),
                ("Готовы к выдаче", summary.get("ready_for_pickup", 0)),
                ("Заправок за период", summary.get("total_refills", 0)),
            ]

        kpi_row_val = [Paragraph(f"<b>{val}</b>", p_kpi_val) for _, val in kpi_items]
        kpi_row_lbl = [Paragraph(lbl, p_kpi_lbl) for lbl, _ in kpi_items]
        kpi_col_width = int(800 / len(kpi_items))
        kpi_widths = [kpi_col_width] * len(kpi_items)

        kpi_table = Table([kpi_row_val, kpi_row_lbl], colWidths=kpi_widths)
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F1F5F9')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        elements.append(kpi_table)
        elements.append(Spacer(1, 8))

        # Основная таблица данных
        items = report.get("items", [])
        if rep_type == "models":
            headers = [
                "№", "Модель картриджа", "Всего шт.", "Доля парка",
                "В работе", "Ожидает заправщика", "На заправке",
                "Готов к выдаче", "Всего заправок"
            ]
            fields = [
                "model", "total", "percentage_label",
                "in_use", "pending_vendor", "at_vendor",
                "ready_for_pickup", "refill_count"
            ]
            col_widths = [25, 175, 65, 75, 75, 95, 85, 85, 120]
        elif rep_type == "history":
            headers = [
                "№", "Дата / Время", "Картридж", "Модель", "Филиал",
                "Операция", "Исполнитель / Сотрудник", "Детали операции"
            ]
            fields = [
                "timestamp", "cartridge_marker", "cartridge_model", "branch_name",
                "action", "user_name", "details"
            ]
            col_widths = [25, 85, 85, 85, 80, 110, 100, 230]
        else:
            headers = [
                "№", "Метка", "Модель", "Кабинет", "Статус",
                "Филиал", "Текущий владелец", "Телефон", "Заправок",
                "Последнее действие", "Дата действия"
            ]
            fields = [
                "marker_label", "model", "cabinet", "status_label",
                "branch_name", "user_name", "user_phone", "refill_count",
                "last_action", "last_action_date"
            ]
            col_widths = [25, 95, 90, 55, 80, 80, 95, 70, 45, 100, 65]

        table_data = [[Paragraph(h, p_head) for h in headers]]

        for idx, item in enumerate(items, 1):
            row_cells = [Paragraph(str(idx), p_cell_c)]
            for fld in fields:
                val = str(item.get(fld, "") or "")
                if fld in ("timestamp", "last_action_date", "refill_count", "cabinet", "status_label", "total", "percentage_label"):
                    row_cells.append(Paragraph(val, p_cell_c))
                else:
                    row_cells.append(Paragraph(val, p_cell))
            table_data.append(row_cells)

        if not items:
            table_data.append([Paragraph("Нет записей по заданным критериям", p_cell_c)] * len(headers))

        data_table = Table(table_data, colWidths=col_widths, repeatRows=1)
        t_style = [
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 2),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
            ('LEFTPADDING', (0, 0), (-1, -1), 3),
            ('RIGHTPADDING', (0, 0), (-1, -1), 3),
        ]
        for r_idx in range(1, len(table_data)):
            if r_idx % 2 == 0:
                t_style.append(('BACKGROUND', (0, r_idx), (-1, r_idx), colors.HexColor('#F8FAFC')))
        data_table.setStyle(TableStyle(t_style))
        elements.append(data_table)

        doc.build(elements)
        stream.seek(0)
        return stream
