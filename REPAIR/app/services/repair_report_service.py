import io
from datetime import datetime
from typing import Optional, Dict, Any, List
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import or_, and_, desc

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from SHARED.models import (
    Asset,
    AssetType,
    AssetStatus,
    AssetCondition,
    Branch,
    ADUser,
    AppUser,
    RepairBatch,
    RepairBatchItem
)

ASSET_TYPE_NAMES_RU = {
    AssetType.WORKSTATION: "Компьютер (ПК)",
    AssetType.LAPTOP: "Ноутбук",
    AssetType.MONITOR: "Монитор",
    AssetType.PRINTER: "Принтер / МФУ",
    AssetType.SERVER: "Сервер",
    AssetType.SWITCH: "Коммутатор",
    AssetType.UPS: "ИБП",
    AssetType.OTHER: "Прочее"
}

ASSET_STATUS_NAMES_RU = {
    AssetStatus.AT_WORKPLACE: "На рабочем месте",
    AssetStatus.PENDING_SC: "Принято в IT (Ожидает СЦ)",
    AssetStatus.AT_SC: "В сервисном центре",
    AssetStatus.RETURNED_IT: "Принято из СЦ (Готово к установке)",
    AssetStatus.DECOMMISSIONED: "Списано"
}

CONDITION_NAMES_RU = {
    AssetCondition.WORKING: "В рабочем состоянии",
    AssetCondition.BROKEN: "В нерабочем состоянии"
}


class RepairReportService:
    @staticmethod
    def get_report_data(
        db: Session,
        current_user: Optional[AppUser] = None,
        condition_filter: Optional[str] = None, # "working", "broken", or None (all)
        asset_type_filter: Optional[str] = None,
        branch_id: Optional[int] = None,
        search_query: Optional[str] = None
    ) -> Dict[str, Any]:
        """Формирует структурированный отчет по оборудованию, состоянию (рабочее/сломанное) и статусам."""
        query = db.query(Asset).options(
            joinedload(Asset.branch),
            joinedload(Asset.responsible_ad_user)
        )

        # Филиальное разграничение
        effective_branch_id = branch_id
        if current_user and current_user.role not in ("superadmin", None) and current_user.branch_id:
            effective_branch_id = current_user.branch_id
            query = query.filter(Asset.branch_id == effective_branch_id)
        elif effective_branch_id:
            query = query.filter(Asset.branch_id == effective_branch_id)

        # Фильтр по техническому состоянию (Рабочее / Сломанное)
        if condition_filter:
            if condition_filter.lower() in ("working", "rabochee"):
                query = query.filter(Asset.condition == AssetCondition.WORKING)
            elif condition_filter.lower() in ("broken", "slomannoe"):
                query = query.filter(Asset.condition == AssetCondition.BROKEN)

        # Фильтр по типу оборудования
        if asset_type_filter:
            try:
                a_type = AssetType(asset_type_filter)
                query = query.filter(Asset.asset_type == a_type)
            except Exception:
                pass

        # Поиск по тексту
        if search_query and search_query.strip():
            sq = f"%{search_query.strip()}%"
            query = query.filter(
                or_(
                    Asset.inventory_number.ilike(sq),
                    Asset.serial_number.ilike(sq),
                    Asset.name.ilike(sq),
                    Asset.hostname.ilike(sq),
                    Asset.cabinet.ilike(sq)
                )
            )

        assets = query.order_by(Asset.inventory_number).all()

        # Статистика
        items = []
        for a in assets:
            user_str = "—"
            if a.responsible_ad_user:
                user_str = a.responsible_ad_user.display_name
            elif a.current_user_id:
                user_str = a.current_user_id

            items.append({
                "id": a.id,
                "inventory_number": a.inventory_number,
                "serial_number": a.serial_number or "—",
                "name": a.name,
                "asset_type": a.asset_type.value if hasattr(a.asset_type, "value") else str(a.asset_type),
                "asset_type_label": ASSET_TYPE_NAMES_RU.get(a.asset_type, str(a.asset_type)),
                "condition": a.condition.value if hasattr(a.condition, "value") else str(a.condition),
                "condition_label": CONDITION_NAMES_RU.get(a.condition, "В рабочем состоянии" if a.condition == AssetCondition.WORKING else "В нерабочем состоянии"),
                "status": a.status.value if hasattr(a.status, "value") else str(a.status),
                "status_label": ASSET_STATUS_NAMES_RU.get(a.status, str(a.status)),
                "branch_name": a.branch.name if a.branch else "Все филиалы",
                "cabinet": a.cabinet or "—",
                "user_name": user_str,
                "hostname": a.hostname or "—",
                "os_name": a.os_name or "—",
                "notes": a.notes or "",
                "updated_at": a.updated_at.strftime("%d.%m.%Y %H:%M") if a.updated_at else "—"
            })

        summary = {
            "total_count": len(items),
            "working_count": sum(1 for i in items if i["condition"] == "working"),
            "broken_count": sum(1 for i in items if i["condition"] == "broken"),
            "at_workplace": sum(1 for i in items if i["status"] == "at_workplace"),
            "pending_sc": sum(1 for i in items if i["status"] == "pending_sc"),
            "at_sc": sum(1 for i in items if i["status"] == "at_sc"),
            "returned_it": sum(1 for i in items if i["status"] == "returned_it"),
            "workstations": sum(1 for i in items if i["asset_type"] in ("workstation", "laptop", "server")),
            "monitors": sum(1 for i in items if i["asset_type"] == "monitor"),
            "printers": sum(1 for i in items if i["asset_type"] == "printer"),
            "ups_count": sum(1 for i in items if i["asset_type"] == "ups"),
            "other_count": sum(1 for i in items if i["asset_type"] in ("switch", "other"))
        }

        branch_obj = db.query(Branch).filter(Branch.id == effective_branch_id).first() if effective_branch_id else None
        branch_title = branch_obj.name if branch_obj else "Все филиалы"

        return {
            "title": "Отчет по состоянию и ремонтам компьютерной техники",
            "branch_name": branch_title,
            "condition_filter": condition_filter or "all",
            "generated_at": datetime.utcnow().strftime("%d.%m.%Y %H:%M"),
            "generated_by": current_user.full_name if current_user else "Система",
            "summary": summary,
            "items": items
        }

    @staticmethod
    def generate_excel(report_data: Dict[str, Any], org_name: str = "IT Отдел") -> io.BytesIO:
        """Генерирует стильный файл Excel (.xlsx) с таблицей техники и сводкой."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Реестр техники"
        ws.views.sheetView[0].showGridLines = True

        # Стили
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid") # Dark slate
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        green_font = Font(name="Calibri", size=10, bold=True, color="166534")

        red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        red_font = Font(name="Calibri", size=10, bold=True, color="991B1B")

        card_title_font = Font(name="Calibri", size=9, bold=True, color="64748B")
        card_val_font = Font(name="Calibri", size=16, bold=True, color="0F172A")

        thin_side = Side(border_style="thin", color="CBD5E1")
        card_border = Border(top=thin_side, left=thin_side, right=thin_side, bottom=thin_side)
        row_border = Border(top=thin_side, left=thin_side, right=thin_side, bottom=thin_side)

        # 1. Шапка документа
        ws["A1"] = f"{org_name} — {report_data.get('title', 'Отчет по технике')}"
        ws["A1"].font = Font(name="Calibri", size=16, bold=True, color="0F172A")
        ws.row_dimensions[1].height = 25

        ws["A2"] = f"Филиал: {report_data.get('branch_name')} | Сформирован: {report_data.get('generated_at')} | Составитель: {report_data.get('generated_by')}"
        ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="475569")
        ws.row_dimensions[2].height = 18

        # 2. Карточки сводки (KPI) в строках 4..5
        summary = report_data.get("summary", {})
        cards = [
            ("Всего единиц", summary.get("total_count", 0), "F1F5F9", "0F172A"),
            ("В рабочем состоянии", summary.get("working_count", 0), "DCFCE7", "166534"),
            ("В нерабочем состоянии", summary.get("broken_count", 0), "FEE2E2", "991B1B"),
            ("На рабочем месте", summary.get("at_workplace", 0), "E0F2FE", "0369A1"),
            ("В ремонте (СЦ / IT)", summary.get("pending_sc", 0) + summary.get("at_sc", 0), "FEF3C7", "B45309"),
            ("Готово к установке", summary.get("returned_it", 0), "F3E8FF", "7E22CE"),
        ]

        col_idx = 1
        for title, val, bg_hex, fg_hex in cards:
            c1 = ws.cell(row=4, column=col_idx, value=title)
            c1.font = Font(name="Calibri", size=9, bold=True, color="64748B")
            c1.alignment = Alignment(horizontal="center", vertical="center")
            c1.fill = PatternFill(start_color=bg_hex, end_color=bg_hex, fill_type="solid")
            c1.border = card_border

            c2 = ws.cell(row=5, column=col_idx, value=val)
            c2.font = Font(name="Calibri", size=16, bold=True, color=fg_hex)
            c2.alignment = Alignment(horizontal="center", vertical="center")
            c2.fill = PatternFill(start_color=bg_hex, end_color=bg_hex, fill_type="solid")
            c2.border = card_border

            col_idx += 1

        ws.row_dimensions[4].height = 18
        ws.row_dimensions[5].height = 28

        # 3. Заголовки таблицы данных
        headers = [
            "№",
            "Инвентарный №",
            "Серийный №",
            "Наименование / Модель",
            "Категория",
            "Состояние техники",
            "Статус",
            "Филиал",
            "Кабинет",
            "Ответственный",
            "Имя ПК / AD",
            "Операционная система",
            "Примечание"
        ]

        start_row = 7
        ws.row_dimensions[start_row].height = 24
        for col_num, h_title in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=col_num, value=h_title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = row_border

        # 4. Строки данных
        current_row = start_row + 1
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        for idx, item in enumerate(report_data.get("items", []), 1):
            ws.row_dimensions[current_row].height = 20
            is_zebra = (idx % 2 == 0)

            cond_is_working = (item.get("condition") == "working")

            row_values = [
                idx,
                item.get("inventory_number"),
                item.get("serial_number"),
                item.get("name"),
                item.get("asset_type_label"),
                item.get("condition_label"),
                item.get("status_label"),
                item.get("branch_name"),
                item.get("cabinet"),
                item.get("user_name"),
                item.get("hostname"),
                item.get("os_name"),
                item.get("notes")
            ]

            for c_idx, val in enumerate(row_values, 1):
                cell = ws.cell(row=current_row, column=c_idx, value=val)
                cell.border = row_border
                cell.font = Font(name="Calibri", size=10)

                # Выделение состояния
                if c_idx == 6: # Состояние техники
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    if cond_is_working:
                        cell.fill = green_fill
                        cell.font = green_font
                    else:
                        cell.fill = red_fill
                        cell.font = red_font
                elif c_idx in (1, 2, 5, 7, 9):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    if is_zebra:
                        cell.fill = zebra_fill
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                    if is_zebra:
                        cell.fill = zebra_fill

            current_row += 1

        # Автоподбор ширины колонок
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row < start_row:
                    continue
                v_str = str(cell.value or "")
                if len(v_str) > max_len:
                    max_len = len(v_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # Специфические ширины
        ws.column_dimensions["A"].width = 6   # №
        ws.column_dimensions["B"].width = 18  # Инв №
        ws.column_dimensions["D"].width = 28  # Модель
        ws.column_dimensions["F"].width = 24  # Состояние
        ws.column_dimensions["G"].width = 24  # Статус
        ws.column_dimensions["J"].width = 24  # Ответственный

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
